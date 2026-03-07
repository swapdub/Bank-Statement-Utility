"""
Export endpoints: CSV, XLSX, JSON — with same filter params as transactions.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime
from typing import Optional

from fastapi import APIRouter, Depends, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, selectinload

from ..database import get_db
from ..models import Transaction, User
from ..auth import get_current_user
from .transactions import _apply_filters

router = APIRouter(prefix="/api/export", tags=["export"])


def _get_export_rows(
    db: Session,
    user_id: int,
    export_all: bool = False,
    search: Optional[str] = None,
    bank_name: Optional[str] = None,
    account_type: Optional[str] = None,
    category_id: Optional[int] = None,
    uncategorized: Optional[bool] = None,
    tag_ids: Optional[str] = None,
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    min_amount: Optional[float] = None,
    max_amount: Optional[float] = None,
    amount_type: Optional[str] = None,
    bank_names: Optional[str] = None,
    category_ids: Optional[str] = None,
    untagged: Optional[bool] = None,
    sort_by: str = "transaction_date",
    sort_order: str = "desc",
):
    """Build filtered query and return Transaction ORM objects."""
    q = db.query(Transaction).options(selectinload(Transaction.tags))

    if export_all:
        # Only filter by user
        q = q.filter(Transaction.user_id == user_id)
    else:
        q = _apply_filters(
            q, db, search, bank_name, account_type, category_id, uncategorized,
            tag_ids, date_from, date_to, min_amount, max_amount, amount_type,
            bank_names, category_ids, untagged, user_id=user_id,
        )

    sort_col = getattr(Transaction, sort_by, Transaction.transaction_date)
    q = q.order_by(sort_col.asc() if sort_order == "asc" else sort_col.desc())

    return q.all()


def _txn_to_dict(txn: Transaction) -> dict:
    """Convert a Transaction ORM object to a flat dict for export."""
    return {
        "id": txn.id,
        "transaction_date": str(txn.transaction_date) if txn.transaction_date else "",
        "description": txn.description or "",
        "debit_amount": txn.debit_amount,
        "credit_amount": txn.credit_amount,
        "closing_balance": txn.closing_balance,
        "bank_name": txn.bank_name or "",
        "account_type": txn.account_type or "",
        "category": txn.category.name if txn.category else "",
        "tags": ", ".join(sorted(t.name for t in txn.tags)) if txn.tags else "",
        "is_transfer": txn.is_transfer,
        "cheque_ref_number": txn.cheque_ref_number or "",
        "value_date": str(txn.value_date) if txn.value_date else "",
    }


_CSV_COLUMNS = [
    "id", "transaction_date", "description", "debit_amount", "credit_amount",
    "closing_balance", "bank_name", "account_type", "category", "tags",
    "is_transfer", "cheque_ref_number", "value_date",
]

_FILTER_PARAMS = dict(
    search=Query(None), bank_name=Query(None), account_type=Query(None),
    category_id=Query(None), uncategorized=Query(None), tag_ids=Query(None),
    date_from=Query(None), date_to=Query(None), min_amount=Query(None),
    max_amount=Query(None), amount_type=Query(None), bank_names=Query(None),
    category_ids=Query(None), untagged=Query(None),
    sort_by=Query("transaction_date"), sort_order=Query("desc"),
    export_all=Query(False),
)


@router.get("/csv")
def export_csv(
    search: Optional[str] = Query(None),
    bank_name: Optional[str] = Query(None),
    account_type: Optional[str] = Query(None),
    category_id: Optional[int] = Query(None),
    uncategorized: Optional[bool] = Query(None),
    tag_ids: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    min_amount: Optional[float] = Query(None),
    max_amount: Optional[float] = Query(None),
    amount_type: Optional[str] = Query(None),
    bank_names: Optional[str] = Query(None),
    category_ids: Optional[str] = Query(None),
    untagged: Optional[bool] = Query(None),
    sort_by: str = Query("transaction_date"),
    sort_order: str = Query("desc"),
    export_all: bool = Query(False),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = _get_export_rows(
        db, user.id, export_all, search, bank_name, account_type, category_id,
        uncategorized, tag_ids, date_from, date_to, min_amount, max_amount,
        amount_type, bank_names, category_ids, untagged, sort_by, sort_order,
    )

    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=_CSV_COLUMNS)
    writer.writeheader()
    for txn in rows:
        writer.writerow(_txn_to_dict(txn))

    buf.seek(0)
    filename = f"transactions_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/xlsx")
def export_xlsx(
    search: Optional[str] = Query(None),
    bank_name: Optional[str] = Query(None),
    account_type: Optional[str] = Query(None),
    category_id: Optional[int] = Query(None),
    uncategorized: Optional[bool] = Query(None),
    tag_ids: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    min_amount: Optional[float] = Query(None),
    max_amount: Optional[float] = Query(None),
    amount_type: Optional[str] = Query(None),
    bank_names: Optional[str] = Query(None),
    category_ids: Optional[str] = Query(None),
    untagged: Optional[bool] = Query(None),
    sort_by: str = Query("transaction_date"),
    sort_order: str = Query("desc"),
    export_all: bool = Query(False),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    rows = _get_export_rows(
        db, user.id, export_all, search, bank_name, account_type, category_id,
        uncategorized, tag_ids, date_from, date_to, min_amount, max_amount,
        amount_type, bank_names, category_ids, untagged, sort_by, sort_order,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    # Header row
    headers = [c.replace("_", " ").title() for c in _CSV_COLUMNS]
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for col_num, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data rows
    for txn in rows:
        d = _txn_to_dict(txn)
        ws.append([d[col] for col in _CSV_COLUMNS])

    # Auto-fit column widths (approximate)
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"transactions_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/json")
def export_json(
    search: Optional[str] = Query(None),
    bank_name: Optional[str] = Query(None),
    account_type: Optional[str] = Query(None),
    category_id: Optional[int] = Query(None),
    uncategorized: Optional[bool] = Query(None),
    tag_ids: Optional[str] = Query(None),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    min_amount: Optional[float] = Query(None),
    max_amount: Optional[float] = Query(None),
    amount_type: Optional[str] = Query(None),
    bank_names: Optional[str] = Query(None),
    category_ids: Optional[str] = Query(None),
    untagged: Optional[bool] = Query(None),
    sort_by: str = Query("transaction_date"),
    sort_order: str = Query("desc"),
    export_all: bool = Query(False),
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    rows = _get_export_rows(
        db, user.id, export_all, search, bank_name, account_type, category_id,
        uncategorized, tag_ids, date_from, date_to, min_amount, max_amount,
        amount_type, bank_names, category_ids, untagged, sort_by, sort_order,
    )

    data = [_txn_to_dict(txn) for txn in rows]
    json_str = json.dumps(data, indent=2, ensure_ascii=False)

    filename = f"transactions_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    return StreamingResponse(
        iter([json_str]),
        media_type="application/json",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
